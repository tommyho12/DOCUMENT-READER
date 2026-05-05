import streamlit as st
import anthropic
import base64
from PIL import Image, ImageEnhance, ImageFilter
import io
import numpy as np
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.enums import TA_LEFT, TA_CENTER

# ── Page Config ───────────────────────────────────────────
st.set_page_config(
    page_title="DocuScan Pro",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ── Styles ────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

.ds-header {
    background: linear-gradient(135deg, #0f2942 0%, #185FA5 100%);
    padding: 24px 32px; border-radius: 12px;
    margin-bottom: 24px; color: white;
    display: flex; align-items: center; justify-content: space-between;
}
.ds-logo { font-size: 24px; font-weight: 700; letter-spacing: -0.5px; }
.ds-logo span { color: #60B4FF; }
.ds-tagline { font-size: 13px; opacity: 0.75; margin-top: 3px; }
.ds-version { font-size: 11px; background: rgba(255,255,255,0.15); padding: 4px 10px; border-radius: 20px; }

.stat-card {
    background: white; border: 1px solid #e8e8e8;
    border-radius: 10px; padding: 16px 20px; text-align: center;
}
.stat-num { font-size: 28px; font-weight: 700; color: #185FA5; }
.stat-lbl { font-size: 12px; color: #999; margin-top: 2px; }

.doc-card {
    background: white; border: 1px solid #e8e8e8;
    border-radius: 10px; padding: 16px;
    margin-bottom: 10px; transition: all 0.2s;
}
.doc-card:hover { border-color: #185FA5; box-shadow: 0 2px 8px rgba(24,95,165,0.1); }
.doc-card-name { font-size: 14px; font-weight: 600; color: #1a1a1a; }
.doc-card-meta { font-size: 12px; color: #999; margin-top: 3px; }
.doc-card-type { font-size: 11px; background: #E6F1FB; color: #185FA5; padding: 2px 8px; border-radius: 20px; font-weight: 500; }

.result-box {
    background: #fafafa; border: 1px solid #e8e8e8;
    border-left: 4px solid #185FA5;
    padding: 18px; border-radius: 8px;
    font-size: 13px; line-height: 1.9;
    white-space: pre-wrap; font-family: 'Courier New', monospace;
    max-height: 400px; overflow-y: auto;
}
.badge-ok { background:#EAF3DE; color:#3B6D11; padding:4px 12px; border-radius:20px; font-size:12px; font-weight:600; }
.badge-proc { background:#E6F1FB; color:#185FA5; padding:4px 12px; border-radius:20px; font-size:12px; font-weight:600; }

.upload-hint {
    background: #FFF8E6; border: 1px solid #F5C842;
    border-radius: 8px; padding: 12px 16px;
    font-size: 13px; color: #7A5C00; margin-bottom: 16px;
}
.section-title { font-size: 16px; font-weight: 600; color: #1a1a1a; margin-bottom: 12px; }
.empty-state { text-align: center; padding: 40px; color: #bbb; font-size: 14px; }
</style>
""", unsafe_allow_html=True)

# ── Session State ─────────────────────────────────────────
if "records" not in st.session_state:
    st.session_state.records = []
if "active_tab" not in st.session_state:
    st.session_state.active_tab = "scan"

# ── Image Enhancement Engine ──────────────────────────────
def enhance_image(img: Image.Image, level: str) -> Image.Image:
    img = img.convert("RGB")
    w, h = img.size

    # Step 1 — Upscale tiny images aggressively
    min_dim = 1200
    if max(w, h) < min_dim:
        scale = min_dim / max(w, h)
        new_w, new_h = int(w * scale), int(h * scale)
        # Extra scale for very tiny images
        if max(w, h) < 300:
            scale = scale * 2
            new_w, new_h = int(w * scale), int(h * scale)
        img = img.resize((new_w, new_h), Image.LANCZOS)

    # Step 2 — Auto white balance / levels
    arr = np.array(img, dtype=np.float32)
    for c in range(3):
        ch = arr[:, :, c]
        p1 = np.percentile(ch, 1)
        p99 = np.percentile(ch, 99)
        if p99 > p1:
            arr[:, :, c] = np.clip((ch - p1) / (p99 - p1) * 255, 0, 255)
    img = Image.fromarray(arr.astype(np.uint8))

    # Step 3 — Enhancement based on level
    settings = {
        "Standard":                  (1.8, 1.2, 3.0, 2),
        "Aggressive (very blurry)":  (2.8, 1.5, 5.0, 3),
        "Max (almost unreadable)":   (4.0, 1.8, 8.0, 4),
    }
    contrast, brightness, sharpness, passes = settings.get(level, settings["Aggressive (very blurry)"])

    img = ImageEnhance.Contrast(img).enhance(contrast)
    img = ImageEnhance.Brightness(img).enhance(brightness)

    for _ in range(passes):
        img = ImageEnhance.Sharpness(img).enhance(sharpness)
        img = img.filter(ImageFilter.SHARPEN)
        img = img.filter(ImageFilter.EDGE_ENHANCE_MORE)

    img = img.filter(ImageFilter.MedianFilter(size=3))
    img = ImageEnhance.Sharpness(img).enhance(sharpness)

    return img

def img_to_b64(img: Image.Image) -> str:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode()

def img_to_bytes(img: Image.Image) -> bytes:
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

# ── AI Prompts ────────────────────────────────────────────
PROMPT = """You are DocuScan Pro — an elite document reading AI built for insurance and tax offices.

This image may be a phone screenshot, blurry photo, low-resolution scan, or dark image. It has already been enhanced. Your job is to extract EVERY piece of readable information — be aggressive and thorough.

IMPORTANT RULES:
1. Never say you cannot read it. Always extract whatever you can see.
2. If text is partially visible, make your best inference and mark it [?]
3. Read in ANY language — do not translate, extract as-is
4. Look for patterns: numbers that look like IDs, dates, dollar amounts, names
5. If you see a table or form, preserve the structure

FORMAT YOUR RESPONSE EXACTLY LIKE THIS:

DOCUMENT TYPE: [identify what this document is]

CLIENT INFORMATION:
[All names, dates of birth, addresses, phone numbers found]

KEY FIELDS:
[All important fields with labels — IDs, policy numbers, amounts, dates, etc.]

FULL EXTRACTED TEXT:
[Every line of text visible, top to bottom, left to right]

CONFIDENCE: [High / Medium / Low] — [brief reason]"""

# ── Export Functions ──────────────────────────────────────
def export_excel(records):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DocuScan Records"

    # Header styling
    header_fill = PatternFill("solid", fgColor="185FA5")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    headers = ["#", "Client Name", "Document Type", "Date Processed", "Extracted Text", "Confidence"]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 30
    col_widths = [5, 25, 20, 20, 80, 15]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    alt_fill = PatternFill("solid", fgColor="EBF3FB")
    for i, rec in enumerate(records, 1):
        row = i + 1
        fill = alt_fill if i % 2 == 0 else PatternFill()
        values = [i, rec["client"], rec["doc_type"], rec["timestamp"],
                  rec["extracted_text"][:500] + ("..." if len(rec["extracted_text"]) > 500 else ""),
                  rec.get("confidence", "—")]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def export_pdf(records):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter,
                            leftMargin=0.75*inch, rightMargin=0.75*inch,
                            topMargin=0.75*inch, bottomMargin=0.75*inch)
    styles = getSampleStyleSheet()
    story = []

    # Title
    title_style = ParagraphStyle("title", fontSize=20, fontName="Helvetica-Bold",
                                  textColor=colors.HexColor("#185FA5"), spaceAfter=4)
    sub_style = ParagraphStyle("sub", fontSize=10, textColor=colors.HexColor("#999999"), spaceAfter=20)
    story.append(Paragraph("DocuScan Pro — Extracted Records", title_style))
    story.append(Paragraph(f"Generated: {datetime.datetime.now().strftime('%B %d, %Y %I:%M %p')} | Total Records: {len(records)}", sub_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#185FA5")))
    story.append(Spacer(1, 16))

    body_style = ParagraphStyle("body", fontSize=9, fontName="Courier", leading=14,
                                 textColor=colors.HexColor("#333333"))
    label_style = ParagraphStyle("label", fontSize=10, fontName="Helvetica-Bold",
                                  textColor=colors.HexColor("#185FA5"), spaceBefore=16, spaceAfter=4)
    meta_style = ParagraphStyle("meta", fontSize=9, textColor=colors.HexColor("#888888"), spaceAfter=8)

    for i, rec in enumerate(records, 1):
        story.append(Paragraph(f"Record #{i} — {rec['client']}", label_style))
        story.append(Paragraph(
            f"Document Type: {rec['doc_type']}  |  Processed: {rec['timestamp']}  |  Confidence: {rec.get('confidence','—')}",
            meta_style))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#dddddd")))
        story.append(Spacer(1, 6))
        safe_text = rec["extracted_text"].replace("<", "&lt;").replace(">", "&gt;")
        story.append(Paragraph(safe_text.replace("\n", "<br/>"), body_style))
        story.append(Spacer(1, 20))

    doc.build(story)
    return buf.getvalue()

# ── Sidebar ───────────────────────────────────────────────
with st.sidebar:
    st.markdown("### ⚙️ Configuration")
    api_key = st.text_input("Anthropic API Key", type="password",
                             placeholder="sk-ant-api03-...",
                             help="Get free at console.anthropic.com → API Keys")
    if api_key and not api_key.startswith("sk-ant-"):
        st.error("Key should start with sk-ant-")
    elif api_key:
        st.success("✓ Key ready")
    st.caption("🔒 Never stored. Only used for this session.")
    st.markdown("[Get a free key →](https://console.anthropic.com)")

    st.divider()
    st.markdown("### 🖼️ Enhancement")
    enhance_level = st.radio("Level", [
        "Standard",
        "Aggressive (very blurry)",
        "Max (almost unreadable)"
    ], index=1)
    st.caption("Use Max for phone screenshots or very dark images.")

    st.divider()
    st.markdown("### 📊 Session Stats")
    total = len(st.session_state.records)
    types = {}
    for r in st.session_state.records:
        types[r["doc_type"]] = types.get(r["doc_type"], 0) + 1
    st.metric("Documents Processed", total)
    if types:
        for t, c in sorted(types.items(), key=lambda x: -x[1]):
            st.caption(f"• {t}: {c}")

    st.divider()
    st.markdown("**DocuScan Pro** v1.0")
    st.caption("Built for Vietrust Insurance")

# ── Header ────────────────────────────────────────────────
st.markdown("""
<div class="ds-header">
  <div>
    <div class="ds-logo">Docu<span>Scan</span> Pro</div>
    <div class="ds-tagline">AI-powered document reading for insurance & tax offices</div>
  </div>
  <div class="ds-version">v1.0 Beta</div>
</div>
""", unsafe_allow_html=True)

# ── Tabs ──────────────────────────────────────────────────
tab_scan, tab_records, tab_export = st.tabs(["📷 Scan Document", "📁 Records", "📤 Export"])

# ════════════════════════════════════════════════════════
# TAB 1 — SCAN
# ════════════════════════════════════════════════════════
with tab_scan:
    st.markdown('<div class="upload-hint">💡 <strong>Works with any image</strong> — phone screenshots, blurry photos, dark scans, any language. The AI reads it all.</div>', unsafe_allow_html=True)

    col_form, col_up = st.columns([1, 2])

    with col_form:
        st.markdown('<div class="section-title">Client Info</div>', unsafe_allow_html=True)
        client_name = st.text_input("Client Name", placeholder="e.g. Nguyen, John")
        doc_type = st.selectbox("Document Type", [
            "Auto-Detect",
            "Insurance Card",
            "Driver's License / ID",
            "Tax Form (W-2 / 1099 / 1040)",
            "Medical Record",
            "Pay Stub",
            "Bank Statement",
            "Letter / Notice",
            "Social Security Card",
            "Medicare / Medicaid Card",
            "Other"
        ])
        notes = st.text_area("Notes (optional)", placeholder="Any context about this document...", height=80)

    with col_up:
        st.markdown('<div class="section-title">Upload Document</div>', unsafe_allow_html=True)
        uploaded = st.file_uploader(
            "Drop image here or click to browse",
            type=["jpg", "jpeg", "png", "webp", "bmp", "tiff", "heic"],
            accept_multiple_files=False,
            label_visibility="collapsed"
        )

        if uploaded:
            img = Image.open(uploaded)
            w, h = img.size
            st.image(img, caption=f"Uploaded: {uploaded.name} ({w}×{h}px)", use_container_width=True)
            if max(w, h) < 400:
                st.warning(f"⚠️ This image is very small ({w}×{h}px). The app will upscale it, but ask your client to send a higher quality photo for best results.")

    st.divider()

    btn_col1, btn_col2, btn_col3 = st.columns([2, 1, 1])
    with btn_col1:
        run_ai = st.button("🔍 Read & Extract with AI", type="primary",
                           disabled=not uploaded, use_container_width=True)
    with btn_col2:
        enhance_btn = st.button("✨ Enhance Image Only", disabled=not uploaded, use_container_width=True)
    with btn_col3:
        clear = st.button("🗑️ Clear", disabled=not uploaded, use_container_width=True)

    # Enhance only
    if enhance_btn and uploaded:
        with st.spinner("Enhancing image..."):
            img = Image.open(uploaded)
            enhanced = enhance_image(img, enhance_level)
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Original**")
            st.image(img, use_container_width=True)
        with c2:
            st.markdown("**Enhanced**")
            st.image(enhanced, use_container_width=True)
        st.download_button("⬇️ Download Enhanced Image", img_to_bytes(enhanced),
                           file_name=f"{uploaded.name}_enhanced.png", mime="image/png")

    # Full AI scan
    if run_ai and uploaded:
        if not api_key:
            st.error("⚠️ Please enter your Anthropic API key in the sidebar first.")
            st.stop()
        if not client_name.strip():
            st.warning("Please enter the client's name before scanning.")
            st.stop()

        with st.spinner("🔍 Enhancing image and reading with AI..."):
            img = Image.open(uploaded)
            enhanced = enhance_image(img, enhance_level)
            b64 = img_to_b64(enhanced)

            try:
                client = anthropic.Anthropic(api_key=api_key)
                # Build prompt with context
                context = f"\nDocument context: Client name is {client_name}. Expected document type: {doc_type}."
                if notes:
                    context += f" Additional notes: {notes}"

                message = client.messages.create(
                    model="claude-opus-4-6",
                    max_tokens=2000,
                    messages=[{
                        "role": "user",
                        "content": [
                            {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": b64}},
                            {"type": "text", "text": PROMPT + context}
                        ]
                    }]
                )

                extracted = message.content[0].text

                # Parse confidence
                confidence = "Medium"
                if "CONFIDENCE: High" in extracted:
                    confidence = "High"
                elif "CONFIDENCE: Low" in extracted:
                    confidence = "Low"

                # Save record
                record = {
                    "id": len(st.session_state.records) + 1,
                    "client": client_name.strip(),
                    "doc_type": doc_type,
                    "filename": uploaded.name,
                    "timestamp": datetime.datetime.now().strftime("%m/%d/%Y %I:%M %p"),
                    "extracted_text": extracted,
                    "confidence": confidence,
                    "notes": notes,
                    "enhanced_img": img_to_bytes(enhanced),
                    "original_size": f"{img.width}×{img.height}px",
                    "enhanced_size": f"{enhanced.width}×{enhanced.height}px",
                }
                st.session_state.records.insert(0, record)

                # Show result
                st.success("✅ Document successfully read and saved to Records!")
                st.markdown(f'<span class="badge-ok">✓ Confidence: {confidence}</span>', unsafe_allow_html=True)
                st.markdown("<br>", unsafe_allow_html=True)

                r1, r2 = st.columns(2)
                with r1:
                    t1, t2 = st.tabs(["Original", "Enhanced"])
                    with t1:
                        st.image(img, use_container_width=True)
                    with t2:
                        st.image(enhanced, use_container_width=True)
                    st.caption(f"Original: {record['original_size']} → Enhanced: {record['enhanced_size']}")

                with r2:
                    st.markdown("**Extracted Text & Data**")
                    st.markdown(f'<div class="result-box">{extracted.replace(chr(60), "&lt;")}</div>', unsafe_allow_html=True)
                    st.download_button("⬇️ Download as TXT", extracted,
                                       file_name=f"{client_name}_{doc_type}.txt", mime="text/plain")

            except anthropic.AuthenticationError:
                st.error("❌ Invalid API key. Please check your key in the sidebar.")
            except anthropic.RateLimitError:
                st.error("❌ Rate limit reached. Please wait 30 seconds and try again.")
            except Exception as e:
                st.error(f"❌ Error: {str(e)}")

# ════════════════════════════════════════════════════════
# TAB 2 — RECORDS
# ════════════════════════════════════════════════════════
with tab_records:
    st.markdown('<div class="section-title">📁 Processed Documents This Session</div>', unsafe_allow_html=True)

    if not st.session_state.records:
        st.markdown('<div class="empty-state">📄<br><br>No documents scanned yet.<br>Go to the Scan tab to get started.</div>', unsafe_allow_html=True)
    else:
        # Search
        search = st.text_input("🔍 Search by client name or document type", placeholder="Type to filter...")

        filtered = st.session_state.records
        if search:
            filtered = [r for r in filtered if
                        search.lower() in r["client"].lower() or
                        search.lower() in r["doc_type"].lower()]

        st.caption(f"Showing {len(filtered)} of {len(st.session_state.records)} records")
        st.divider()

        for rec in filtered:
            with st.expander(f"#{rec['id']}  {rec['client']}  —  {rec['doc_type']}  |  {rec['timestamp']}"):
                left, right = st.columns([1, 2])
                with left:
                    if rec.get("enhanced_img"):
                        st.image(rec["enhanced_img"], caption="Enhanced view", use_container_width=True)
                    st.caption(f"File: {rec['filename']}")
                    st.caption(f"Size: {rec['original_size']} → {rec['enhanced_size']}")
                    st.caption(f"Confidence: {rec.get('confidence','—')}")
                    if rec.get("notes"):
                        st.caption(f"Notes: {rec['notes']}")
                with right:
                    st.markdown("**Extracted Data**")
                    st.markdown(f'<div class="result-box">{rec["extracted_text"].replace("<","&lt;")}</div>', unsafe_allow_html=True)
                    st.download_button(
                        "⬇️ Download TXT",
                        rec["extracted_text"],
                        file_name=f"{rec['client']}_{rec['doc_type']}.txt",
                        mime="text/plain",
                        key=f"dl_{rec['id']}"
                    )

# ════════════════════════════════════════════════════════
# TAB 3 — EXPORT
# ════════════════════════════════════════════════════════
with tab_export:
    st.markdown('<div class="section-title">📤 Export All Records</div>', unsafe_allow_html=True)

    if not st.session_state.records:
        st.markdown('<div class="empty-state">📄<br><br>No records to export yet.<br>Scan some documents first.</div>', unsafe_allow_html=True)
    else:
        st.info(f"Ready to export **{len(st.session_state.records)} records** from this session.")

        c1, c2, c3 = st.columns(3)

        with c1:
            st.markdown("### 📊 Excel")
            st.markdown("All records in a formatted spreadsheet. Best for sorting, filtering, and entering data into your system.")
            if st.button("Generate Excel File", use_container_width=True):
                excel_data = export_excel(st.session_state.records)
                st.download_button(
                    "⬇️ Download Excel",
                    excel_data,
                    file_name=f"DocuScan_Export_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )

        with c2:
            st.markdown("### 📄 PDF Report")
            st.markdown("Professional PDF with all extracted text. Best for printing or sharing with clients.")
            if st.button("Generate PDF Report", use_container_width=True):
                pdf_data = export_pdf(st.session_state.records)
                st.download_button(
                    "⬇️ Download PDF",
                    pdf_data,
                    file_name=f"DocuScan_Report_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

        with c3:
            st.markdown("### 📋 Text Summary")
            st.markdown("Simple plain text file with all extracted data. Best for copying into other systems.")
            if st.button("Generate Text File", use_container_width=True):
                lines = []
                for rec in st.session_state.records:
                    lines.append(f"{'='*60}")
                    lines.append(f"Client: {rec['client']}")
                    lines.append(f"Type: {rec['doc_type']}")
                    lines.append(f"Date: {rec['timestamp']}")
                    lines.append(f"Confidence: {rec.get('confidence','—')}")
                    lines.append(f"{'─'*40}")
                    lines.append(rec["extracted_text"])
                    lines.append("")
                txt = "\n".join(lines)
                st.download_button(
                    "⬇️ Download TXT",
                    txt,
                    file_name=f"DocuScan_Summary_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.txt",
                    mime="text/plain",
                    use_container_width=True
                )

        st.divider()
        if st.button("🗑️ Clear All Records from Session", type="secondary"):
            st.session_state.records = []
            st.rerun()
